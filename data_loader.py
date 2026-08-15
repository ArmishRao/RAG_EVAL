"""
Data Loader for Programming Documentation Assistant
Supports: CSV, JSON, JSONL, TXT, PDF, Excel, DOCX, MD, HTML
"""

import os
import json
import glob
import pandas as pd
from typing import List, Dict, Any, Optional
from langchain_core.documents import Document
import logging
import re

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentationLoader:
    """Handle loading documentation from multiple formats"""
    
    def __init__(self, data_folder: str = "docs_data"):
        self.data_folder = data_folder
        self.supported_formats = {
            '.csv': self._load_csv,
            '.json': self._load_json,
            '.jsonl': self._load_jsonl,
            '.txt': self._load_text,
            '.pdf': self._load_pdf,
            '.xlsx': self._load_excel,
            '.xls': self._load_excel,
            '.docx': self._load_docx,
            '.md': self._load_markdown,
            '.html': self._load_html,
            '.htm': self._load_html,
            '.rst': self._load_rst,
        }
    
    def load_all(self) -> List[Document]:
        """Load all documents from all supported formats"""
        documents = []
        
        if not os.path.exists(self.data_folder):
            logger.warning(f"Data folder '{self.data_folder}' not found")
            return documents
        
        for root, dirs, files in os.walk(self.data_folder):
            for file in files:
                file_path = os.path.join(root, file)
                ext = os.path.splitext(file)[1].lower()
                
                if ext in self.supported_formats:
                    logger.info(f"📄 Loading: {file_path}")
                    try:
                        docs = self.supported_formats[ext](file_path)
                        documents.extend(docs)
                        logger.info(f"   ✅ Added {len(docs)} documents from {file}")
                    except Exception as e:
                        logger.error(f"   ❌ Error loading {file}: {e}")
        
        return documents
    
    def _load_markdown(self, file_path: str) -> List[Document]:
        """Load Markdown files with section extraction"""
        documents = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Extract sections based on markdown headings
            sections = re.split(r'\n#{1,6}\s+', content)
            headings = re.findall(r'^#{1,6}\s+(.+)$', content, re.MULTILINE)
            
            if len(sections) > 1:
                for i, section in enumerate(sections[1:], 1):
                    if len(section.strip()) > 50:
                        heading = headings[i-1] if i-1 < len(headings) else f"Section {i}"
                        metadata = {
                            'source_file': os.path.basename(file_path),
                            'heading': heading,
                            'source_type': 'markdown',
                            'file_path': file_path,
                            'doc_type': self._detect_doc_type(file_path)
                        }
                        documents.append(Document(
                            page_content=f"## {heading}\n\n{section}",
                            metadata=metadata
                        ))
            else:
                # Single document
                metadata = {
                    'source_file': os.path.basename(file_path),
                    'source_type': 'markdown',
                    'file_path': file_path,
                    'doc_type': self._detect_doc_type(file_path)
                }
                documents.append(Document(page_content=content, metadata=metadata))
                
        except Exception as e:
            logger.error(f"Error loading Markdown {file_path}: {e}")
        
        return documents
    
    def _load_html(self, file_path: str) -> List[Document]:
        """Load HTML files (documentation pages)"""
        documents = []
        
        try:
            from bs4 import BeautifulSoup
            
            with open(file_path, 'r', encoding='utf-8') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')
            
            # Remove script and style elements
            for element in soup(["script", "style", "nav", "header", "footer"]):
                element.decompose()
            
            # Try to find main content
            main_content = soup.find('main') or soup.find('article') or soup.find('div', {'class': 'content'})
            
            if main_content:
                content = main_content.get_text(separator='\n', strip=True)
            else:
                content = soup.get_text(separator='\n', strip=True)
            
            # Clean up
            content = '\n'.join([line.strip() for line in content.split('\n') if line.strip()])
            
            # Get title
            title = soup.title.string if soup.title else os.path.basename(file_path)
            
            if len(content.strip()) > 100:
                metadata = {
                    'source_file': os.path.basename(file_path),
                    'title': str(title)[:200],
                    'source_type': 'html',
                    'file_path': file_path,
                    'doc_type': self._detect_doc_type(file_path)
                }
                documents.append(Document(page_content=content, metadata=metadata))
                
        except Exception as e:
            logger.error(f"Error loading HTML {file_path}: {e}")
        
        return documents
    
    def _load_rst(self, file_path: str) -> List[Document]:
        """Load reStructuredText files"""
        documents = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            # Extract sections from RST
            sections = re.split(r'\n={3,}\n', content)
            headings = re.findall(r'^(.+)\n={3,}$', content, re.MULTILINE)
            
            if len(sections) > 1:
                for i, section in enumerate(sections):
                    if len(section.strip()) > 50:
                        heading = headings[i] if i < len(headings) else f"Section {i+1}"
                        metadata = {
                            'source_file': os.path.basename(file_path),
                            'heading': heading,
                            'source_type': 'rst',
                            'file_path': file_path,
                            'doc_type': self._detect_doc_type(file_path)
                        }
                        documents.append(Document(
                            page_content=f"# {heading}\n\n{section}",
                            metadata=metadata
                        ))
            else:
                metadata = {
                    'source_file': os.path.basename(file_path),
                    'source_type': 'rst',
                    'file_path': file_path,
                    'doc_type': self._detect_doc_type(file_path)
                }
                documents.append(Document(page_content=content, metadata=metadata))
                
        except Exception as e:
            logger.error(f"Error loading RST {file_path}: {e}")
        
        return documents
    
    def _load_text(self, file_path: str) -> List[Document]:
        """Load plain text files"""
        documents = []
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            if len(content.strip()) < 50:
                return documents
            
            metadata = {
                'source_file': os.path.basename(file_path),
                'source_type': 'text',
                'file_path': file_path,
                'doc_type': self._detect_doc_type(file_path)
            }
            documents.append(Document(page_content=content, metadata=metadata))
            
        except Exception as e:
            logger.error(f"Error loading text {file_path}: {e}")
        
        return documents
    
    def _load_pdf(self, file_path: str) -> List[Document]:
        """Load PDF files"""
        documents = []
        
        try:
            try:
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    full_text = ""
                    for page_num, page in enumerate(pdf.pages):
                        text = page.extract_text()
                        if text and len(text.strip()) > 50:
                            full_text += text + "\n\n"
                            metadata = {
                                'source_file': os.path.basename(file_path),
                                'page_number': page_num + 1,
                                'source_type': 'pdf',
                                'file_path': file_path,
                                'doc_type': self._detect_doc_type(file_path)
                            }
                            documents.append(Document(page_content=text, metadata=metadata))
                    
                    if full_text.strip():
                        metadata = {
                            'source_file': os.path.basename(file_path),
                            'source_type': 'pdf',
                            'file_path': file_path,
                            'doc_type': self._detect_doc_type(file_path)
                        }
                        documents.append(Document(page_content=full_text, metadata=metadata))
                        
            except ImportError:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    full_text = ""
                    for page_num, page in enumerate(reader.pages):
                        text = page.extract_text()
                        if text and len(text.strip()) > 50:
                            full_text += text + "\n\n"
                            metadata = {
                                'source_file': os.path.basename(file_path),
                                'page_number': page_num + 1,
                                'source_type': 'pdf',
                                'file_path': file_path,
                                'doc_type': self._detect_doc_type(file_path)
                            }
                            documents.append(Document(page_content=text, metadata=metadata))
                    
                    if full_text.strip():
                        metadata = {
                            'source_file': os.path.basename(file_path),
                            'source_type': 'pdf',
                            'file_path': file_path,
                            'doc_type': self._detect_doc_type(file_path)
                        }
                        documents.append(Document(page_content=full_text, metadata=metadata))
                        
        except Exception as e:
            logger.error(f"Error loading PDF {file_path}: {e}")
        
        return documents
    
    # Keep other load methods from original...
    def _load_csv(self, file_path: str) -> List[Document]:
        """Load CSV files - adapted for documentation"""
        documents = []
        try:
            df = pd.read_csv(file_path)
            for _, row in df.iterrows():
                content = " ".join([str(v) for v in row.values if pd.notna(v)])
                if len(content.strip()) > 50:
                    metadata = {
                        'source_file': os.path.basename(file_path),
                        'source_type': 'csv',
                        'doc_type': self._detect_doc_type(file_path)
                    }
                    documents.append(Document(page_content=content, metadata=metadata))
        except Exception as e:
            logger.error(f"Error loading CSV {file_path}: {e}")
        return documents
    
    def _load_json(self, file_path: str) -> List[Document]:
        """Load JSON files"""
        documents = []
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            if isinstance(data, list):
                for item in data:
                    doc = self._process_json_item(item, file_path)
                    if doc:
                        documents.append(doc)
            elif isinstance(data, dict):
                doc = self._process_json_item(data, file_path)
                if doc:
                    documents.append(doc)
        except Exception as e:
            logger.error(f"Error loading JSON {file_path}: {e}")
        return documents
    
    def _load_jsonl(self, file_path: str) -> List[Document]:
        """Load JSONL files"""
        documents = []
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line:
                        try:
                            item = json.loads(line)
                            doc = self._process_json_item(item, file_path)
                            if doc:
                                documents.append(doc)
                        except json.JSONDecodeError:
                            continue
        except Exception as e:
            logger.error(f"Error loading JSONL {file_path}: {e}")
        return documents
    
    def _load_excel(self, file_path: str) -> List[Document]:
        """Load Excel files"""
        documents = []
        try:
            xl = pd.ExcelFile(file_path)
            for sheet_name in xl.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                for _, row in df.iterrows():
                    content = " ".join([str(v) for v in row.values if pd.notna(v)])
                    if len(content.strip()) > 50:
                        metadata = {
                            'source_file': os.path.basename(file_path),
                            'sheet': sheet_name,
                            'source_type': 'excel',
                            'doc_type': self._detect_doc_type(file_path)
                        }
                        documents.append(Document(page_content=content, metadata=metadata))
        except Exception as e:
            logger.error(f"Error loading Excel {file_path}: {e}")
        return documents
    
    def _load_docx(self, file_path: str) -> List[Document]:
        """Load Word documents"""
        documents = []
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            if paragraphs:
                content = "\n".join(paragraphs)
                metadata = {
                    'source_file': os.path.basename(file_path),
                    'source_type': 'docx',
                    'file_path': file_path,
                    'doc_type': self._detect_doc_type(file_path)
                }
                documents.append(Document(page_content=content, metadata=metadata))
        except Exception as e:
            logger.error(f"Error loading DOCX {file_path}: {e}")
        return documents
    
    def _process_json_item(self, item: Dict, file_path: str) -> Optional[Document]:
        """Process a JSON item into a Document"""
        content_keys = ['content', 'text', 'description', 'body', 'summary', 'docstring']
        content = None
        
        for key in content_keys:
            if key in item and item[key]:
                content = str(item[key])
                break
        
        if not content or len(content.strip()) < 50:
            return None
        
        metadata = {
            'source_file': os.path.basename(file_path),
            'source_type': 'json',
            'file_path': file_path,
            'doc_type': self._detect_doc_type(file_path)
        }
        
        for key, value in item.items():
            if key not in content_keys and value is not None:
                metadata[key] = str(value)[:500]
        
        return Document(page_content=content, metadata=metadata)
    
    def _detect_doc_type(self, file_path: str) -> str:
        """Detect documentation type from file path"""
        path_lower = file_path.lower()
        if 'python' in path_lower:
            return 'python'
        elif 'langchain' in path_lower:
            return 'langchain'
        elif 'fastapi' in path_lower:
            return 'fastapi'
        elif 'pytorch' in path_lower:
            return 'pytorch'
        elif 'tensorflow' in path_lower:
            return 'tensorflow'
        elif 'django' in path_lower:
            return 'django'
        elif 'flask' in path_lower:
            return 'flask'
        else:
            return 'other'


# Keep as DataLoader for backward compatibility
DataLoader = DocumentationLoader