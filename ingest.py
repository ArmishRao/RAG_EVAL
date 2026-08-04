"""
Multi-Format Data Ingestion for Pakistan Legal Advisor
Supports: CSV, JSON, JSONL, TXT, PDF, Excel, DOCX
Creates FAISS vector store for semantic search only
"""

import os
import json
import pandas as pd
from typing import List, Optional, Dict, Any
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
import logging
import torch
import gc
import time
from pathlib import Path

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
torch.set_num_threads(8)


# ============================================================
# CONFIGURATION
# ============================================================

DATASET_FOLDER = "datasets"          # Primary folder (backward compatible)
DATA_FOLDER = "data"                 # New multi-format folder
VECTORSTORE_FOLDER = "vectorstore"
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
CHUNK_SIZE = 1500
CHUNK_OVERLAP = 100


# ============================================================
# CSV LOADING FUNCTIONS
# ============================================================

def read_csv_file(file_path: str) -> Optional[pd.DataFrame]:
    """Read CSV with better error handling."""
    encodings = ["utf-8", "utf-8-sig", "cp1252", "latin1"]
    
    for encoding in encodings:
        try:
            df = pd.read_csv(
                file_path, 
                encoding=encoding,
                quotechar='"',
                escapechar='\\',
                on_bad_lines='skip'
            )
            logger.info(f"Loaded CSV with {encoding}")
            return df
        except UnicodeDecodeError:
            continue
        except Exception as e:
            logger.debug(f"Error with {encoding}: {e}")
            continue
    
    raise Exception(f"Cannot read {file_path}")


def load_csv_documents(csv_file: str) -> List[Document]:
    """Load documents from CSV files."""
    documents = []
    df = read_csv_file(csv_file)
    
    if df is None:
        return documents
    
    filename = os.path.basename(csv_file)
    
    # Check if it's the existing format with required columns
    required_columns = ["Book", "Chapter Number", "Chapter Title", "Section", "Heading", "Defination"]
    has_required = all(col in df.columns for col in required_columns)
    
    if has_required:
        # Process with existing format
        for _, row in df.iterrows():
            if pd.isna(row["Defination"]):
                continue
            
            text = f"""
Book: {row['Book']}
Chapter Number: {row['Chapter Number']}
Chapter Title: {row['Chapter Title']}
Section: {row['Section']}
Heading: {row['Heading']}
Definition: {row['Defination']}
"""
            
            metadata = {
                "book": str(row["Book"]),
                "chapter_number": str(row["Chapter Number"]),
                "chapter_title": str(row["Chapter Title"]),
                "section": str(row["Section"]),
                "heading": str(row["Heading"]),
                "source_file": filename,
                "source_type": "csv"
            }
            
            documents.append(Document(page_content=text, metadata=metadata))
    else:
        # Generic CSV - combine all columns
        logger.info(f"Processing generic CSV: {filename}")
        for _, row in df.iterrows():
            content = " ".join([str(v) for v in row.values if pd.notna(v)])
            if len(content.strip()) > 50:
                metadata = {
                    "source_file": filename,
                    "source_type": "csv",
                    "columns": ", ".join(df.columns)
                }
                # Add all columns as metadata
                for col in df.columns:
                    if pd.notna(row[col]):
                        metadata[col] = str(row[col])[:200]
                
                documents.append(Document(page_content=content, metadata=metadata))
    
    return documents


# ============================================================
# JSON LOADING FUNCTIONS
# ============================================================

def load_json_documents(file_path: str) -> List[Document]:
    """Load documents from JSON files."""
    documents = []
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        filename = os.path.basename(file_path)
        
        items_to_process = []
        
        if isinstance(data, list):
            items_to_process = data
        elif isinstance(data, dict):
            nested_keys = ['data', 'laws', 'documents', 'content', 'items', 'records']
            found = False
            for key in nested_keys:
                if key in data and isinstance(data[key], list):
                    items_to_process = data[key]
                    found = True
                    break
            
            if not found:
                # Single object
                items_to_process = [data]
        
        for item in items_to_process:
            doc = process_json_item(item, filename)
            if doc:
                documents.append(doc)
    
    except Exception as e:
        logger.error(f"Error loading JSON {file_path}: {e}")
    
    return documents


def process_json_item(item: Dict, filename: str) -> Optional[Document]:
    """Process a single JSON item into a Document."""
    # Try to find content from common keys
    content = None
    content_keys = ['content', 'text', 'definition', 'defination', 'description', 
                   'body', 'details', 'summary', 'full_text']
    
    for key in content_keys:
        if key in item and item[key]:
            content = str(item[key])
            break
    
    # If no content found, try to create from all values
    if not content:
        content = " ".join([str(v) for v in item.values() if isinstance(v, (str, int, float))])
    
    if not content or len(content.strip()) < 50:
        return None
    
    # Build metadata
    metadata = {
        'source_file': filename,
        'source_type': 'json'
    }
    
    # Add all other fields as metadata (limit length)
    for key, value in item.items():
        if key not in content_keys and value is not None:
            try:
                if isinstance(value, (str, int, float, bool)):
                    metadata[key] = str(value)[:500]
                elif isinstance(value, list):
                    metadata[key] = str(len(value)) + " items"
            except:
                pass
    
    return Document(page_content=content, metadata=metadata)


def load_jsonl_documents(file_path: str) -> List[Document]:
    """Load documents from JSONL (JSON Lines) files."""
    documents = []
    filename = os.path.basename(file_path)
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            for line_num, line in enumerate(f):
                line = line.strip()
                if line:
                    try:
                        item = json.loads(line)
                        doc = process_json_item(item, filename)
                        if doc:
                            doc.metadata['line_number'] = line_num + 1
                            documents.append(doc)
                    except json.JSONDecodeError:
                        continue
    except Exception as e:
        logger.error(f"Error loading JSONL {file_path}: {e}")
    
    return documents


# ============================================================
# TEXT LOADING FUNCTIONS
# ============================================================

def load_text_documents(file_path: str) -> List[Document]:
    """Load documents from plain text files."""
    documents = []
    filename = os.path.basename(file_path)
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        if len(content.strip()) < 50:
            return documents
        
        # Try to split by sections (double newline or headings)
        sections = []
        
        # Method 1: Double newline
        if '\n\n' in content:
            parts = content.split('\n\n')
            if len(parts) > 3:
                for i, part in enumerate(parts):
                    if len(part.strip()) > 100:
                        sections.append({
                            'content': part,
                            'section': i + 1
                        })
        
        # Method 2: Look for section headers
        if not sections:
            import re
            section_pattern = r'(?i)(?:section|chapter|article|part)\s+(\d+[A-Za-z]?)\s*[:.]?\s*(.+?)(?=(?:(?:section|chapter|article|part)\s+\d+[A-Za-z]?|$))'
            matches = re.findall(section_pattern, content, re.DOTALL)
            
            for match in matches:
                section_num = match[0]
                section_content = match[1].strip()
                if section_content and len(section_content) > 50:
                    sections.append({
                        'content': section_content,
                        'section': section_num
                    })
        
        if sections:
            for sec in sections:
                metadata = {
                    'source_file': filename,
                    'source_type': 'text',
                    'section': str(sec.get('section', 'Unknown'))
                }
                documents.append(Document(page_content=sec['content'], metadata=metadata))
        else:
            # Single document
            metadata = {
                'source_file': filename,
                'source_type': 'text'
            }
            documents.append(Document(page_content=content, metadata=metadata))
    
    except Exception as e:
        logger.error(f"Error loading text file {file_path}: {e}")
    
    return documents


# ============================================================
# PDF LOADING FUNCTIONS
# ============================================================

def load_pdf_documents(file_path: str) -> List[Document]:
    """Load documents from PDF files."""
    documents = []
    filename = os.path.basename(file_path)
    
    try:
        # Try pdfplumber first
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                full_text = ""
                for page_num, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text and len(text.strip()) > 50:
                        full_text += text + "\n\n"
                        
                        metadata = {
                            'source_file': filename,
                            'source_type': 'pdf',
                            'page_number': page_num + 1,
                            'total_pages': len(pdf.pages)
                        }
                        documents.append(Document(page_content=text, metadata=metadata))
                
                if full_text.strip():
                    metadata = {
                        'source_file': filename,
                        'source_type': 'pdf',
                        'total_pages': len(pdf.pages)
                    }
                    documents.append(Document(page_content=full_text, metadata=metadata))
        
        except ImportError:
            # Fallback to PyPDF2
            try:
                import PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    full_text = ""
                    for page_num, page in enumerate(reader.pages):
                        text = page.extract_text()
                        if text and len(text.strip()) > 50:
                            full_text += text + "\n\n"
                            metadata = {
                                'source_file': filename,
                                'source_type': 'pdf',
                                'page_number': page_num + 1,
                                'total_pages': len(reader.pages)
                            }
                            documents.append(Document(page_content=text, metadata=metadata))
                    
                    if full_text.strip():
                        metadata = {
                            'source_file': filename,
                            'source_type': 'pdf',
                            'total_pages': len(reader.pages)
                        }
                        documents.append(Document(page_content=full_text, metadata=metadata))
            except ImportError:
                logger.warning("PDF support requires: pip install pdfplumber or PyPDF2")
    
    except Exception as e:
        logger.error(f"Error loading PDF {file_path}: {e}")
    
    return documents


# ============================================================
# EXCEL LOADING FUNCTIONS
# ============================================================

def load_excel_documents(file_path: str) -> List[Document]:
    """Load documents from Excel files."""
    documents = []
    filename = os.path.basename(file_path)
    
    try:
        xl = pd.ExcelFile(file_path)
        
        for sheet_name in xl.sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            # Try to find text columns
            text_cols = [col for col in df.columns if any(
                kw in str(col).lower() 
                for kw in ['text', 'content', 'definition', 'description', 'details', 'info']
            )]
            
            if text_cols:
                for _, row in df.iterrows():
                    content = str(row[text_cols[0]]) if pd.notna(row[text_cols[0]]) else ""
                    if len(content.strip()) > 50:
                        metadata = {
                            'source_file': filename,
                            'sheet': sheet_name,
                            'source_type': 'excel'
                        }
                        for col in df.columns:
                            if col != text_cols[0] and pd.notna(row[col]):
                                metadata[str(col)] = str(row[col])[:200]
                        documents.append(Document(page_content=content, metadata=metadata))
            else:
                for _, row in df.iterrows():
                    content = " ".join([str(v) for v in row.values if pd.notna(v)])
                    if len(content.strip()) > 50:
                        metadata = {
                            'source_file': filename,
                            'sheet': sheet_name,
                            'source_type': 'excel'
                        }
                        documents.append(Document(page_content=content, metadata=metadata))
    
    except Exception as e:
        logger.error(f"Error loading Excel {file_path}: {e}")
    
    return documents


# ============================================================
# DOCX LOADING FUNCTIONS
# ============================================================

def load_docx_documents(file_path: str) -> List[Document]:
    """Load documents from Word (DOCX) files."""
    documents = []
    filename = os.path.basename(file_path)
    
    try:
        from docx import Document as DocxDocument
        
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        if not paragraphs:
            return documents
        
        full_text = "\n".join(paragraphs)
        
        # Try to identify sections by headings
        headings = []
        sections = []
        current_heading = "Introduction"
        current_content = []
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
                
            if para.style.name and ('Heading' in para.style.name or 'Title' in para.style.name):
                if current_content:
                    sections.append({
                        'heading': current_heading,
                        'content': "\n".join(current_content)
                    })
                    current_content = []
                current_heading = text
            else:
                current_content.append(text)
        
        if current_content:
            sections.append({
                'heading': current_heading,
                'content': "\n".join(current_content)
            })
        
        if sections:
            for section in sections:
                content = f"**{section['heading']}**\n\n{section['content']}"
                metadata = {
                    'source_file': filename,
                    'source_type': 'docx',
                    'heading': section['heading']
                }
                documents.append(Document(page_content=content, metadata=metadata))
        else:
            metadata = {
                'source_file': filename,
                'source_type': 'docx'
            }
            documents.append(Document(page_content=full_text, metadata=metadata))
    
    except ImportError:
        logger.warning("DOCX support requires: pip install python-docx")
    except Exception as e:
        logger.error(f"Error loading DOCX {file_path}: {e}")
    
    return documents


# ============================================================
# MAIN LOADING FUNCTIONS
# ============================================================

def load_documents_from_folder(folder_path: str) -> List[Document]:
    """Load documents from a folder with multiple formats."""
    documents = []
    
    if not os.path.exists(folder_path):
        return documents
    
    handlers = {
        '.csv': load_csv_documents,
        '.json': load_json_documents,
        '.jsonl': load_jsonl_documents,
        '.txt': load_text_documents,
        '.text': load_text_documents,
        '.pdf': load_pdf_documents,
        '.xlsx': load_excel_documents,
        '.xls': load_excel_documents,
        '.docx': load_docx_documents,
    }
    
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            ext = os.path.splitext(file)[1].lower()
            
            if ext in handlers:
                logger.info(f"Loading: {file_path}")
                try:
                    docs = handlers[ext](file_path)
                    documents.extend(docs)
                    logger.info(f"Added {len(docs)} documents from {file}")
                except Exception as e:
                    logger.error(f"Error loading {file}: {e}")
    
    return documents


def load_all_documents() -> List[Document]:
    """Load documents from all configured folders."""
    all_documents = []
    
    if os.path.exists(DATASET_FOLDER):
        logger.info(f"Loading from {DATASET_FOLDER} (CSV format)...")
        docs = load_documents_from_folder(DATASET_FOLDER)
        all_documents.extend(docs)
        logger.info(f"Loaded {len(docs)} documents from {DATASET_FOLDER}")
    
    if os.path.exists(DATA_FOLDER):
        logger.info(f"Loading from {DATA_FOLDER} (multi-format)...")
        docs = load_documents_from_folder(DATA_FOLDER)
        all_documents.extend(docs)
        logger.info(f"Loaded {len(docs)} documents from {DATA_FOLDER}")
    
    return all_documents


# ============================================================
# VECTORSTORE CREATION
# ============================================================

def create_vectorstore(documents: List[Document]):
    """
    Optimized FAISS vector store creation.
    """
    if not documents:
        logger.error("No documents found to ingest.")
        return
    
    start_time = time.time()
    logger.info(f"Total Documents Loaded: {len(documents)}")
    
    # Split documents into chunks
    logger.info("Splitting documents into chunks...")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ".", " ", ""],
        keep_separator=True,
        length_function=len,
    )
    chunks = splitter.split_documents(documents)
    logger.info(f"Created {len(chunks)} chunks")
    
    # Initialize embeddings
    logger.info(f"Loading embedding model: {EMBEDDING_MODEL}")
    embeddings = HuggingFaceEmbeddings(
        model_name=EMBEDDING_MODEL,
        encode_kwargs={"batch_size": 64, "show_progress_bar": True}
    )
    
    # Build FAISS index in batches
    batch_size = 1000
    total_batches = (len(chunks) + batch_size - 1) // batch_size
    vectorstore = None
    
    logger.info(f"Processing {total_batches} batches of {batch_size} chunks...")
    
    for i in range(0, len(chunks), batch_size):
        batch = chunks[i:i+batch_size]
        batch_num = i // batch_size + 1
        
        logger.info(f"Batch {batch_num}/{total_batches} ({len(batch)} chunks)")
        
        try:
            if vectorstore is None:
                vectorstore = FAISS.from_documents(batch, embeddings)
            else:
                vectorstore.add_documents(batch)
            
            gc.collect()
            
        except Exception as e:
            logger.error(f"Error processing batch {batch_num}: {e}")
            continue
    
    if vectorstore is None:
        logger.error("No vectorstore created.")
        return
    
    # Save vectorstore
    Path(VECTORSTORE_FOLDER).mkdir(exist_ok=True)
    
    logger.info(f"Saving vectorstore to {VECTORSTORE_FOLDER}...")
    vectorstore.save_local(VECTORSTORE_FOLDER)
    
    elapsed = time.time() - start_time
    logger.info(f"Vectorstore saved successfully in {elapsed:.1f} seconds")


# ============================================================
# MAIN ENTRY POINT
# ============================================================

def main():
    """Main entry point for data ingestion."""
    print("=" * 70)
    print("Pakistan Legal Advisor - Data Ingestion")
    print("=" * 70)
    print("\nSupported Formats:")
    print("   CSV, JSON, JSONL, TXT, PDF, Excel, DOCX")
    print("\nVector Search Support:")
    print("   FAISS vector store for semantic search")
    print("=" * 70)
    
    print("\nLoading documents...")
    documents = load_all_documents()
    
    if not documents:
        print("\nNo documents found to ingest.")
        print("\nPlease add your data files to:")
        print("    'datasets/'    - CSV files (legacy format)")
        print("    'data/'        - Any supported format")
        print("\nOr download the Pakistan Laws Dataset:")
        print("   https://huggingface.co/AyeshaJadoon/Pakistan_Laws_Dataset")
        return
    
    print(f"\nLoaded {len(documents)} documents")
    
    create_vectorstore(documents)
    
    print("\n" + "=" * 70)
    print("Done! Vector store created successfully.")
    print(f"\nFiles Created:")
    print(f"   {VECTORSTORE_FOLDER}/ - FAISS index")
    print("=" * 70)
    print("\nYou can now run the retriever.")


if __name__ == "__main__":
    main()